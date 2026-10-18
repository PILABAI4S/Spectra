"""Read-only common workbook interface; actual pool sizes follow the input."""

from pathlib import Path

import numpy as np
import openpyxl


def read_workbook(path, reaction_col):
    with Path(path).open("rb") as source:
        workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
        try:
            sheet = workbook.active
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            index = {key: i for i, key in enumerate(headers)}
            if reaction_col not in index:
                raise ValueError("Missing identifier column: " + reaction_col)
            rows = []
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if not any(v is not None and v != "" for v in values):
                    continue
                record = {key: values[i] if i < len(values) else None for key, i in index.items()}
                reaction = str(record.get(reaction_col) or "").strip()
                if "-" not in reaction:
                    raise ValueError("Invalid sample identifier: " + reaction)
                conversion = record.get("Conversion")
                if conversion is None:
                    conversion = record.get("Conc_Mean")
                record.update(Conversion=conversion, Conc_Mean=conversion, Reaction=reaction,
                              Batch=reaction.split("-")[0], Suffix=int(reaction.split("-")[1]))
                rows.append(record)
            return rows
        finally:
            workbook.close()


def is_hit(record):
    return bool(9000 <= float(record["Mw"]) <= 11000 and float(record["PDI"]) < 1.5
                and float(record["Conc_Mean"]) > 70)


def load_data(initial_file, all_file):
    initial = read_workbook(initial_file, "Reaction")
    rows = read_workbook(all_file, "Reaction_ID")
    for label, records in [("initial", initial), ("database", rows)]:
        if not records or len({r["Reaction"] for r in records}) != len(records):
            raise ValueError("Require nonempty unique sample IDs in " + label)
        groups = {}
        for record in records:
            if record["Reaction"] != "{}-{:02d}".format(record["Batch"], record["Suffix"]):
                raise ValueError("Invalid sample ID: " + record["Reaction"])
            for key in ["Total_mols", "Initiator", "Time", "Temp", "Mw", "PDI", "Conc_Mean"]:
                if not np.isfinite(float(record[key])):
                    raise ValueError("Non-finite " + key)
            groups.setdefault(record["Batch"], []).append(record)
        for batch, members in groups.items():
            if len(members) != 12 or {r["Suffix"] for r in members} != set(range(1, 13)):
                raise ValueError("Require12 states in batch " + batch)
            if any(float(r["Time"]) != r["Suffix"] * 120 for r in members):
                raise ValueError("Require120-min sampling intervals: " + batch)
            if any(len({float(r[k]) for r in members}) != 1 for k in ["Total_mols", "Initiator", "Temp"]):
                raise ValueError("Inconsistent batch conditions: " + batch)
    lookup = {r["Reaction"]: r for r in rows}
    if not {r["Reaction"] for r in initial}.issubset(lookup):
        raise ValueError("Initial samples absent from full database")
    batches = {}
    for row in rows:
        batches.setdefault(row["Batch"], []).append(row)
    for group in batches.values():
        group.sort(key=lambda r: r["Suffix"])
    mismatches = [{"reaction": r["Reaction"], "fields": [k for k in ["Total_mols", "Initiator", "Time", "Temp", "Mw", "PDI", "Conc_Mean"]
                  if not np.isclose(float(r[k]), float(lookup[r["Reaction"]][k]), rtol=1e-6, atol=1e-6)]} for r in initial]
    check = {"database_states": len(rows), "database_batches": len(batches), "initial_states": len(initial),
             "initial_batches": len({r["Batch"] for r in initial}),
             "initial_target_satisfying_states": sum(is_hit(lookup[r["Reaction"]]) for r in initial),
             "initial_database_value_mismatches": [r for r in mismatches if r["fields"]],
             "values_source": "All workbook; Initial workbook selects revealed IDs"}
    return initial, rows, lookup, batches, check
